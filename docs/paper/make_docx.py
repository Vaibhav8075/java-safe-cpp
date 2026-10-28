"""Builds docs/paper/JavaSafeCpp_IEEE.docx (IEEE conference layout in Word) from the LaTeX paper.

The Word file is generated from main.tex, generated/numbers.tex, generated/tab_*.tex and the
bibliography produced by BibTeX, so it always carries the same text and the same numbers as the PDF.

Usage (after docs/paper/build.sh has built the PDF):  python docs/paper/make_docx.py
Requires: python-docx, Pillow, and pdftoppm / pdftotext (Poppler) on the PATH.
"""
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
BUILD = HERE.parent.parent / "build" / "paper"
OUT = HERE / "JavaSafeCpp_IEEE.docx"
FONT = "Times New Roman"
MONO = "Courier New"

tex = (HERE / "main.tex").read_text(encoding="utf-8").replace("\r\n", "\n")
numbers_tex = (HERE / "generated/numbers.tex").read_text(encoding="utf-8")
bbl = (BUILD / "main.bbl").read_text(encoding="utf-8").replace("\r\n", "\n")

RES = {k: v for k, v in re.findall(r"\\csname res@(\w+)\\endcsname\{(.*)\}", numbers_tex)}

# ------------------------------------------------------------------ bibliography numbering
bib_items = re.split(r"\\bibitem\{([^}]+)\}", bbl)
bib_keys, bib_text = [], []
for i in range(1, len(bib_items), 2):
    bib_keys.append(bib_items[i])
    bib_text.append(bib_items[i + 1].split("\\end{thebibliography}")[0])
CITE_NO = {k: n + 1 for n, k in enumerate(bib_keys)}

# ------------------------------------------------------------------ labels (section numbers etc.)
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
LABELS = {}
body_start = tex.index("\\maketitle") + len("\\maketitle")
body = tex[body_start:tex.index("\\bibliographystyle")]
body = re.sub(r"(?<!\\)%.*", "", body)                       # LaTeX comments
body = re.sub(r"\\begin\{abstract\}.*?\\end\{abstract\}", "", body, flags=re.S)
body = re.sub(r"\\begin\{IEEEkeywords\}.*?\\end\{IEEEkeywords\}", "", body, flags=re.S)

sec_no = sub_no = tab_no = fig_no = 0
last_label_target = None
for m in re.finditer(r"\\section\{[^}]*\}|\\subsection\{[^}]*\}|\\begin\{table\*?\}|\\begin\{figure\*?\}|\\label\{([^}]+)\}", body):
    s = m.group(0)
    if s.startswith("\\section"):
        sec_no, sub_no = sec_no + 1, 0
        last_label_target = ROMAN[sec_no - 1]
    elif s.startswith("\\subsection"):
        sub_no += 1
        last_label_target = f"{ROMAN[sec_no - 1]}-{chr(64 + sub_no)}"
    elif s.startswith("\\begin{table"):
        tab_no += 1
        last_label_target = ROMAN[tab_no - 1]
    elif s.startswith("\\begin{figure"):
        fig_no += 1
        last_label_target = str(fig_no)
    else:
        LABELS[m.group(1)] = last_label_target

ACCENT = {"'": "\u0301", "`": "\u0300", '"': "\u0308", "^": "\u0302", "~": "\u0303"}


# ------------------------------------------------------------------ inline LaTeX -> runs
def read_group(s, i):
    """s[i] == '{'; returns (content, index after the closing brace)."""
    depth, j = 0, i
    while j < len(s):
        if s[j] == "\\":
            j += 2
            continue
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces: " + s[i:i + 40])


def math_runs(m, st):
    m = m.replace("\\times", "\u00d7").replace("\\ge", "\u2265").replace("\\le", "\u2264").replace("-", "\u2212")
    out, i = [], 0
    while i < len(m):
        if m[i] == "^":
            if m[i + 1] == "{":
                g, i = read_group(m, i + 1)
            else:
                g, i = m[i + 1], i + 2
            out.append((g, dict(st, sup=True)))
        else:
            out.append((m[i], dict(st)))
            i += 1
    return out


def inline(s, st=None):
    st = dict(st or {})
    runs, i = [], 0

    def add(text, **extra):
        if text:
            runs.append((text, dict(st, **extra)))

    while i < len(s):
        c = s[i]
        if c == "\\":
            if i + 1 >= len(s):
                break
            n = s[i + 1]
            if n.isalpha():
                j = i + 1
                while j < len(s) and s[j].isalpha():
                    j += 1
                cmd = s[i + 1:j]
                arg = None
                if cmd in {"texttt", "emph", "textbf", "res", "cite", "ref", "url", "textsc", "item"} and j < len(s):
                    k = j
                    while k < len(s) and s[k] == " ":
                        k += 1
                    if k < len(s) and s[k] == "{":
                        arg, j = read_group(s, k)
                if cmd == "texttt":
                    runs.extend(inline(arg, dict(st, mono=True)))
                elif cmd == "emph":
                    runs.extend(inline(arg, dict(st, italic=True)))
                elif cmd == "textbf":
                    runs.extend(inline(arg, dict(st, bold=True)))
                elif cmd == "textsc":
                    runs.extend(inline(arg, dict(st, caps=True)))
                elif cmd == "res":
                    runs.extend(inline(RES[arg], st))
                elif cmd == "url":
                    add(arg, mono=True)
                elif cmd == "ref":
                    add(LABELS[arg])
                elif cmd == "cite":
                    nums = [CITE_NO[k.strip()] for k in arg.split(",")]
                    add(", ".join(f"[{x}]" for x in nums))
                elif cmd == "jscpp":
                    add("jscpp", caps=True)
                    if j < len(s) and s[j] == " ":
                        j += 1
                elif cmd in {"newblock", "BIBentryALTinterwordspacing", "BIBentrySTDinterwordspacing"}:
                    add(" ")
                elif cmd == "times":
                    add("\u00d7")
                i = j
                continue
            if n in ACCENT:                                   # \'a  or  \'{a}
                k = i + 2
                if k < len(s) and s[k] == "{":
                    letter, k = read_group(s, k)
                else:
                    letter, k = s[k], k + 1
                add(unicodedata.normalize("NFC", letter + ACCENT[n]))
                i = k
                continue
            if n == "\\":
                add("\n")
            elif n == ",":
                add("\u2009")
            elif n == " ":
                add(" ")
            else:                                             # \& \% \_ \# \{ \}
                add(n)
            i += 2
            continue
        if c == "{":
            if s.startswith("{,}", i):
                add(",")
                i += 3
                continue
            g, i2 = read_group(s, i)
            runs.extend(inline(g, st))
            i = i2
            continue
        if c == "}":
            i += 1
            continue
        if c == "$":
            j = s.index("$", i + 1)
            runs.extend(math_runs(s[i + 1:j], st))
            i = j + 1
            continue
        if c == "~":
            add("\u00a0")
        elif s.startswith("---", i):
            add("\u2014")
            i += 3
            continue
        elif s.startswith("--", i):
            add("\u2013")
            i += 2
            continue
        elif s.startswith("``", i):
            add("\u201c")
            i += 2
            continue
        elif s.startswith("''", i):
            add("\u201d")
            i += 2
            continue
        elif c == "\n":
            add(" ")
        else:
            add(c)
        i += 1
    return runs


# ------------------------------------------------------------------ python-docx helpers
doc = Document()
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)

for name, italic, caps, align in [("Heading 1", False, True, WD_ALIGN_PARAGRAPH.CENTER),
                                  ("Heading 2", True, False, WD_ALIGN_PARAGRAPH.LEFT)]:
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold, h.font.italic = FONT, Pt(10), False, italic
    h.font.small_caps = caps
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        h.element.rPr.rFonts.attrib.pop(qn(attr), None)
    h.paragraph_format.alignment = align
    h.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 6)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True


def set_columns(section, n):
    sect = section._sectPr
    for old in sect.findall(qn("w:cols")):
        sect.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "360")
    # schema order: ... pgSz, pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols, formProt, ...
    sect.insert_element_before(cols, "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
                               "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")


def setup_page(section):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.625)
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(1.0)


def emit(par, runs, size=None):
    for text, st in runs:
        parts = text.split("\n")
        for n, part in enumerate(parts):
            r = par.add_run(part)
            r.font.name = MONO if st.get("mono") else FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO if st.get("mono") else FONT)
            if size:
                r.font.size = Pt(size - 1 if st.get("mono") else size)
            elif st.get("mono"):
                r.font.size = Pt(9)
            r.bold = st.get("bold") or None
            r.italic = st.get("italic") or None
            r.font.small_caps = st.get("caps") or None
            r.font.superscript = st.get("sup") or None
            if n < len(parts) - 1:
                r.add_break()


def para(text, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, first_indent=0.17, st=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(first_indent)
    emit(p, inline(text, st), size)
    return p


def shade(par, fill):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.insert_element_before(
        shd, "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
        "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
        "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange")


def box(par):
    ppr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "2")
        b.set(qn("w:color"), "808080")
        borders.append(b)
    ppr.append(borders)


def cell_border(cell, **sides):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side, size in sides.items():
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tcpr.append(borders)


def caption(label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.keep_with_next = True
    emit(p, inline(label, {"caps": True}), 8)
    p.add_run().add_break()
    emit(p, inline(text, {"caps": True}), 8)


def add_table(fragment, widths):
    lines = [l.strip() for l in fragment.strip().splitlines()]
    header = None
    rows = []
    stage = 0
    for l in lines:
        if l.startswith("\\begin{tabular}") or l in {"\\end{tabular}", "\\toprule", "\\bottomrule"}:
            continue
        if l == "\\midrule":
            stage += 1
            continue
        cells = [c.strip() for c in re.split(r"(?<!\\)&", l.rstrip("\\").strip())]
        if stage == 0:
            header = cells
        else:
            rows.append(cells)
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ri, cells in enumerate([header] + rows):
        for ci, text in enumerate(cells):
            cell = table.cell(ri, ci)
            cell.width = Inches(widths[ci])
            p = cell.paragraphs[0]
            plain = "".join(t for t, _ in inline(text))
            numeric = ci > 0 and re.fullmatch(r"[\d,.\u2212\u00d7 ]+", plain) is not None
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT
            emit(p, inline(text, {"bold": ri == 0}), 8)
            if ri == 0:
                cell_border(cell, top=8, bottom=4)
            if ri == len(rows):
                cell_border(cell, bottom=8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure_image():
    """Crops the pipeline figure out of the compiled PDF.

    Finds the page whose text contains the 'Fig. 1.' caption, renders it, and keeps the first block of ink
    below the top margin: the figure ends at the first blank gap (the space before its caption).
    """
    pdf = HERE / "main.pdf"
    page = next(n for n in range(1, 9)
                if "Fig. 1." in subprocess.run(["pdftotext", "-f", str(n), "-l", str(n), str(pdf), "-"],
                                               capture_output=True, text=True).stdout)
    dpi = 300
    png = BUILD / "figure_pipeline"
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-f", str(page), "-l", str(page), "-singlefile",
                    str(pdf), str(png)], check=True)
    im = Image.open(str(png) + ".png").convert("L")
    w, h = im.size
    left, right = int(0.55 * dpi), int(7.95 * dpi)
    px = im.load()

    def ink(y):
        return any(px[x, y] < 200 for x in range(left, right, 2))

    y = int(0.5 * dpi)
    while y < h and not ink(y):
        y += 1
    top = y
    gap = 0
    while y < h and gap < int(0.12 * dpi):
        gap = 0 if ink(y) else gap + 1
        y += 1
    bottom = y - gap
    out = Path(str(png) + "_crop.png")
    im.crop((left, max(top - 6, 0), right, bottom + 6)).save(out)
    return out


# ------------------------------------------------------------------ title block
setup_page(doc.sections[0])
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(10)
r = title.add_run("Java-Safe C++: A Bytecode Compiler and Managed Virtual Machine for a Memory-Safe C++ Subset")
r.font.size, r.font.name = Pt(24), FONT

def blocks(command):
    out = []
    for m in re.finditer(re.escape(command) + r"\{", tex):
        out.append(read_group(tex, m.end() - 1)[0])
    return out


authors = blocks("\\IEEEauthorblockN")
affil = blocks("\\IEEEauthorblockA")
tbl = doc.add_table(rows=1, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for ci in range(3):
    cell = tbl.cell(0, ci)
    cell.width = Inches(2.4)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit(p, inline(authors[ci], {"italic": False}), 11)
    for line in re.split(r"\\\\\s*", affil[ci]):
        line = line.strip().replace("\n", " ")
        if line:
            q = cell.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            emit(q, inline(line, {"italic": True}), 9)
doc.add_paragraph().paragraph_format.space_after = Pt(4)

# body section: two columns
doc.add_section(WD_SECTION.CONTINUOUS)
setup_page(doc.sections[-1])
set_columns(doc.sections[-1], 2)

abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", tex, re.S).group(1).strip()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
emit(p, [("Abstract\u2014", {"bold": True, "italic": True})] + inline(abstract, {"bold": True}), 9)
p.paragraph_format.space_after = Pt(6)
kw = re.search(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", tex, re.S).group(1).strip()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
emit(p, [("Index Terms\u2014", {"bold": True, "italic": True})] + inline(kw, {"bold": True}), 9)
p.paragraph_format.space_after = Pt(6)

# ------------------------------------------------------------------ body
buf = []
sec_i = sub_i = tab_i = 0


def flush():
    text = " ".join(" ".join(buf).split())
    buf.clear()
    if text:
        para(text)


token = re.compile(
    r"\\section\{([^}]*)\}|\\subsection\{([^}]*)\}"
    r"|\\begin\{(itemize|enumerate|lstlisting|table\*?|figure\*?)\}(.*?)\\end\{\3\}"
    r"|\n\s*\n|\\label\{[^}]*\}", re.S)
pos = 0
for m in token.finditer(body):
    buf.append(body[pos:m.start()])
    pos = m.end()
    if m.group(1) is not None:
        flush()
        sec_i, sub_i = sec_i + 1, 0
        doc.add_paragraph(f"{ROMAN[sec_i - 1]}. {m.group(1)}", style="Heading 1")
    elif m.group(2) is not None:
        flush()
        sub_i += 1
        doc.add_paragraph(f"{chr(64 + sub_i)}. {m.group(2)}", style="Heading 2")
    elif m.group(0).startswith("\n"):
        flush()
    elif m.group(3) is not None:
        flush()
        env, content = m.group(3), m.group(4)
        if env in ("itemize", "enumerate"):
            items = [" ".join(x.split()) for x in re.split(r"\\item\s*", content)[1:]]
            for n, it in enumerate(items, 1):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.18)
                p.paragraph_format.space_after = Pt(2)
                emit(p, [(f"{n})" if env == "enumerate" else "\u2022", {}), ("\u00a0" if env == "enumerate" else "\u00a0\u00a0", {})] + inline(it))
        elif env == "lstlisting":
            code_lines = content.strip("\n").split("\n")
            for k, line in enumerate(code_lines):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.05)
                shade(p, "F2F2F2")
                r = p.add_run(line.replace(" ", "\u00a0"))
                r.font.name, r.font.size = MONO, Pt(8)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif env.startswith("table"):
            tab_i += 1
            cap = re.search(r"\\caption\{((?:[^{}]|\{[^{}]*\})*)\}", content, re.S).group(1)
            fn = re.search(r"\\input\{([^}]+)\}", content).group(1)
            caption(f"TABLE {ROMAN[tab_i - 1]}", " ".join(cap.split()))
            widths = {"tab_loc": [2.3, 0.8], "tab_safety": [1.45, 0.5, 1.3, 0.35], "tab_perf": [1.25, 0.5, 0.65, 0.65, 0.45]}[Path(fn).stem]
            add_table((HERE / fn).read_text(encoding="utf-8"), widths)
        elif env.startswith("figure"):
            cap = re.search(r"\\caption\{((?:[^{}]|\{[^{}]*\})*)\}", content, re.S).group(1)
            # full-width figure: switch to one column, then back to two
            doc.add_section(WD_SECTION.CONTINUOUS)
            set_columns(doc.sections[-1], 1)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(figure_image()), width=Inches(6.9))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            emit(p, inline("Fig. 1. " + " ".join(cap.split())), 8)
            p.paragraph_format.space_after = Pt(6)
            doc.add_section(WD_SECTION.CONTINUOUS)
            set_columns(doc.sections[-1], 2)
flush_text = body[pos:]
buf.append(flush_text)
flush()

# ------------------------------------------------------------------ references
doc.add_paragraph("References", style="Heading 1")
for n, text in enumerate(bib_text, 1):
    text = re.sub(r"\\newblock", " ", text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(2)
    emit(p, [(f"[{n}]\u00a0\u00a0", {})] + inline(" ".join(text.split()), {}), 8)

# the last section keeps the two-column layout; make sure every section uses letter paper
for s in doc.sections:
    setup_page(s)

zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")

doc.core_properties.title ="Java-Safe C++: A Bytecode Compiler and Managed Virtual Machine for a Memory-Safe C++ Subset"
doc.core_properties.author = "Vaibhav Goel"
doc.save(OUT)
print("wrote", OUT.relative_to(HERE.parent.parent))
